# ============================================
# МАГАЗИН АНТИКВАРНОЙ МЕБЕЛИ «АНТИКВАР»
# Главный серверный файл (Flask + Excel + JSON)
# ============================================
# Возможности:
#  - Авторизация и регистрация (scrypt-хеши паролей)
#  - Товары с НЕСКОЛЬКИМИ фото (в Excel — через запятую)
#  - Корзина и избранное (сервер для авторизованных, localStorage для гостей)
#  - Заказы со статусами (Excel), валюта — белорусский рубль (Br)
#  - Админ-панель: добавление / редактирование / удаление товаров
#  - Защита от 500: безопасное преобразование всех чисел
# ============================================

from flask import Flask, render_template, request, redirect, url_for, session, jsonify, flash
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from werkzeug.utils import secure_filename
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from datetime import datetime
import os
import json
import traceback

# ============================================
# НАСТРОЙКА ПРИЛОЖЕНИЯ
# ============================================
app = Flask(__name__,
            template_folder='html',
            static_folder='.',
            static_url_path='')

app.secret_key = 'antiquar_secret_key_2026_change_this'
app.config['TEMPLATES_AUTO_RELOAD'] = True
app.config['SEND_FILE_MAX_AGE_DEFAULT'] = 0

# Пути к файлам
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, 'data')
PRODUCTS_FILE = os.path.join(DATA_DIR, 'products.xlsx')
USERS_FILE = os.path.join(DATA_DIR, 'users.xlsx')
ORDERS_FILE = os.path.join(DATA_DIR, 'orders.xlsx')
USER_DATA_FILE = os.path.join(DATA_DIR, 'user_data.json')
UPLOAD_FOLDER = os.path.join(BASE_DIR, 'images', 'uploads')

# Создаём папки
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024  # 16 МБ
ALLOWED_EXTENSIONS = {'png', 'jpg', 'jpeg', 'gif', 'webp'}

# Заголовки таблицы заказов (валюта — белорусский рубль)
ORDERS_HEADERS = ['ID заказа', 'Дата и время', 'Клиент', 'Телефон', 'Логин',
                  'Состав заказа (ID:кол-во)', 'Сумма (Br)', 'Комментарий', 'Статус']

# Цвета статусов заказа: (цвет текста, цвет фона)
STATUS_COLORS = {
    'Новый': ('27ae60', 'E8F8F0'),
    'В обработке': ('F39C12', 'FFF4E0'),
    'Выполнен': ('2980B9', 'E8F0FE'),
    'Отменён': ('E74C3C', 'FDE8E8')
}


def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def safe_int(value, default=0):
    """Безопасное преобразование в int — никогда не бросает исключение (защита от 500)."""
    try:
        return int(value)
    except (ValueError, TypeError):
        return default


# ============================================
# ИНИЦИАЛИЗАЦИЯ ФАЙЛОВ
# ============================================
def init_users_file():
    if not os.path.exists(USERS_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Пользователи'
        headers = ['ID', 'Логин', 'Имя', 'Телефон', 'Пароль', 'Роль', 'Дата регистрации']
        ws.append(headers)
        header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2C1810', end_color='2C1810', fill_type='solid')
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 20
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 18
        ws.column_dimensions['E'].width = 60
        ws.column_dimensions['F'].width = 12
        ws.column_dimensions['G'].width = 20
        admin_password = generate_password_hash('admin')
        ws.append([
            1, 'admin', 'Администратор', '+7 (000) 000-00-00',
            admin_password, 'admin', datetime.now().strftime('%d.%m.%Y')
        ])
        wb.save(USERS_FILE)
        print("✅ Файл пользователей создан. Админ: admin / admin")


def init_products_file():
    if not os.path.exists(PRODUCTS_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Товары'
        headers = ['ID', 'Название', 'Категория', 'Описание', 'Цена', 'Бейдж', 'Фото', 'Продано']
        ws.append(headers)
        wb.save(PRODUCTS_FILE)
        print("✅ Файл товаров создан")


def init_orders_file():
    if not os.path.exists(ORDERS_FILE):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Заказы'
        ws.append(ORDERS_HEADERS)
        header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='2C1810', end_color='2C1810', fill_type='solid')
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        widths = {'A': 12, 'B': 20, 'C': 25, 'D': 18, 'E': 20, 'F': 40, 'G': 15, 'H': 30, 'I': 15}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w
        ws.freeze_panes = 'A2'
        wb.save(ORDERS_FILE)
        print("✅ Файл заказов создан")


def init_user_data_file():
    if not os.path.exists(USER_DATA_FILE):
        with open(USER_DATA_FILE, 'w', encoding='utf-8') as f:
            json.dump({}, f, ensure_ascii=False, indent=2)
        print("✅ Файл данных пользователей создан")


# Инициализация при запуске
init_users_file()
init_products_file()
init_orders_file()
init_user_data_file()


# ============================================
# РАБОТА С ПОЛЬЗОВАТЕЛЯМИ
# ============================================
def load_users():
    if not os.path.exists(USERS_FILE):
        return []
    wb = load_workbook(USERS_FILE)
    ws = wb.active
    users = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            users.append({
                'id': safe_int(row[0]),
                'login': str(row[1]) if row[1] else '',
                'name': str(row[2]) if row[2] else '',
                'phone': str(row[3]) if row[3] else '',
                'password': str(row[4]) if row[4] else '',
                'role': str(row[5]) if row[5] else 'client',
                'registered': str(row[6]) if row[6] else ''
            })
    return users


def get_user_by_login(login):
    for user in load_users():
        if user['login'].lower() == login.lower():
            return user
    return None


def get_user_by_id(user_id):
    for user in load_users():
        if user['id'] == safe_int(user_id):
            return user
    return None


def register_user(login, name, phone, password):
    if get_user_by_login(login):
        return False, 'Пользователь с таким логином уже существует'
    users = load_users()
    new_id = max((u['id'] for u in users), default=0) + 1
    wb = load_workbook(USERS_FILE)
    ws = wb.active
    hashed_password = generate_password_hash(password)
    ws.append([
        new_id, login, name, phone, hashed_password, 'client',
        datetime.now().strftime('%d.%m.%Y')
    ])
    wb.save(USERS_FILE)
    return True, 'Регистрация успешна'


def authenticate_user(login, password):
    user = get_user_by_login(login)
    if not user:
        return None, 'Пользователь не найден'
    if check_password_hash(user['password'], password):
        return user, 'Успешный вход'
    return None, 'Неверный пароль'


# ============================================
# ДАННЫЕ ПОЛЬЗОВАТЕЛЯ (КОРЗИНА, ИЗБРАННОЕ)
# ============================================
def load_user_data():
    if not os.path.exists(USER_DATA_FILE):
        return {}
    try:
        with open(USER_DATA_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception:
        return {}


def save_user_data(data):
    with open(USER_DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def get_user_cart(user_id):
    data = load_user_data()
    return data.get(str(user_id), {}).get('cart', {})


def get_user_favorites(user_id):
    data = load_user_data()
    return data.get(str(user_id), {}).get('favorites', [])


def save_user_cart(user_id, cart):
    data = load_user_data()
    user_id_str = str(user_id)
    if user_id_str not in data:
        data[user_id_str] = {}
    data[user_id_str]['cart'] = cart
    save_user_data(data)


def save_user_favorites(user_id, favorites):
    data = load_user_data()
    user_id_str = str(user_id)
    if user_id_str not in data:
        data[user_id_str] = {}
    data[user_id_str]['favorites'] = favorites
    save_user_data(data)


# ============================================
# РАБОТА С ТОВАРАМИ (МУЛЬТИ-ФОТО)
# ============================================
def load_products():
    """Читает товары. Колонка «Фото» может содержать несколько путей через запятую."""
    if not os.path.exists(PRODUCTS_FILE):
        return []
    wb = load_workbook(PRODUCTS_FILE)
    ws = wb.active
    products = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            images_raw = str(row[6]) if row[6] else ''
            images = [p.strip() for p in images_raw.split(',') if p.strip()]
            products.append({
                'id': safe_int(row[0]),
                'name': str(row[1]) if row[1] else '',
                'category': str(row[2]) if row[2] else '',
                'description': str(row[3]) if row[3] else '',
                'price': safe_int(row[4]),
                'badge': str(row[5]) if row[5] else '',
                'image': images[0] if images else '',   # первое фото — главное
                'images': images,                        # все фото списком
                'sold': bool(row[7]) if len(row) > 7 and row[7] else False
            })
    return products


def save_products(products):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Товары'
    ws.append(['ID', 'Название', 'Категория', 'Описание', 'Цена', 'Бейдж', 'Фото', 'Продано'])
    header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2C1810', end_color='2C1810', fill_type='solid')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    for product in products:
        images = product.get('images') or ([product['image']] if product.get('image') else [])
        ws.append([
            product['id'], product['name'], product['category'],
            product['description'], product['price'], product['badge'],
            ','.join(images), product['sold']
        ])
    widths = {'A': 8, 'B': 30, 'C': 15, 'D': 50, 'E': 12, 'F': 12, 'G': 40, 'H': 10}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    wb.save(PRODUCTS_FILE)


def get_next_id():
    products = load_products()
    if not products:
        return 1
    return max(p['id'] for p in products) + 1


# ============================================
# ДЕКОРАТОРЫ АВТОРИЗАЦИИ
# ============================================
def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function


def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('login'))
        if session.get('role') != 'admin':
            flash('У вас нет прав администратора', 'error')
            return redirect(url_for('index'))
        return f(*args, **kwargs)
    return decorated_function


# ============================================
# МАРШРУТЫ: АВТОРИЗАЦИЯ И РЕГИСТРАЦИЯ
# ============================================
@app.route('/test_route')
def test_route():
    return "Маршрут работает!"


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        login_name = request.form.get('login', '').strip()
        password = request.form.get('password', '').strip()
        if not login_name or not password:
            return render_template('login.html', error='Заполните все поля')
        user, message = authenticate_user(login_name, password)
        if user:
            session['user_id'] = user['id']
            session['username'] = user['name']
            session['login'] = user['login']
            session['role'] = user['role']
            if user['role'] == 'admin':
                return redirect(url_for('admin'))
            return redirect(url_for('index'))
        else:
            return render_template('login.html', error=message)
    return render_template('login.html')


@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        login_name = request.form.get('login', '').strip()
        name = request.form.get('name', '').strip()
        phone = request.form.get('phone', '').strip()
        password = request.form.get('password', '').strip()
        password_confirm = request.form.get('password_confirm', '').strip()

        if not login_name or not name or not password:
            return render_template('register.html', error='Заполните обязательные поля')
        if password != password_confirm:
            return render_template('register.html', error='Пароли не совпадают')
        if len(password) < 6:
            return render_template('register.html', error='Пароль должен содержать минимум 6 символов')
        if not login_name.isalnum():
            return render_template('register.html', error='Логин должен содержать только буквы и цифры')

        success, message = register_user(login_name, name, phone, password)
        if success:
            user = get_user_by_login(login_name)
            session['user_id'] = user['id']
            session['username'] = user['name']
            session['login'] = user['login']
            session['role'] = user['role']
            return redirect(url_for('index'))
        else:
            return render_template('register.html', error=message)
    return render_template('register.html')


@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('index'))


# ============================================
# МАРШРУТЫ: ГЛАВНАЯ СТРАНИЦА
# ============================================
@app.route('/')
def index():
    return render_template('index.html',
                           products=load_products(),
                           username=session.get('username'),
                           user_role=session.get('role'))


# ============================================
# МАРШРУТЫ: АДМИН-ПАНЕЛЬ
# ============================================
@app.route('/admin')
@admin_required
def admin():
    products = load_products()
    sold_count = sum(1 for p in products if p['sold'])
    available_count = len(products) - sold_count
    return render_template('admin.html',
                           products=products,
                           username=session['username'],
                           sold_count=sold_count,
                           available_count=available_count)


@app.route('/add_product', methods=['POST'])
@admin_required
def add_product():
    name = (request.form.get('name') or '').strip()
    category = (request.form.get('category') or '').strip()
    description = (request.form.get('description') or '').strip()
    price = safe_int(request.form.get('price'))
    badge = (request.form.get('badge') or '').strip()

    if not name or not category:
        flash('Заполните название и категорию товара', 'error')
        return redirect(url_for('admin'))

    # --- Загрузка НЕСКОЛЬКИХ фото ---
    image_paths = []
    files = request.files.getlist('images')
    if not files and 'image' in request.files:      # совместимость со старой формой
        files = [request.files['image']]
    for i, file in enumerate(files):
        if file and file.filename and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            name_part, ext = os.path.splitext(filename)
            unique_filename = f"{timestamp}_{i}_{name_part}{ext}"
            filepath = os.path.join(app.config['UPLOAD_FOLDER'], unique_filename)
            file.save(filepath)
            image_paths.append(f"images/uploads/{unique_filename}")

    products = load_products()
    new_product = {
        'id': get_next_id(),
        'name': name,
        'category': category,
        'description': description,
        'price': price,
        'badge': badge,
        'image': image_paths[0] if image_paths else '',
        'images': image_paths,
        'sold': False
    }
    products.append(new_product)
    save_products(products)
    flash('✅ Товар добавлен', 'success')
    return redirect(url_for('admin'))


@app.route('/delete_product/<int:product_id>')
@admin_required
def delete_product(product_id):
    products = load_products()
    products = [p for p in products if p['id'] != product_id]
    save_products(products)
    return redirect(url_for('admin'))


@app.route('/toggle_sold/<int:product_id>')
@admin_required
def toggle_sold(product_id):
    products = load_products()
    for p in products:
        if p['id'] == product_id:
            p['sold'] = not p['sold']
            break
    save_products(products)
    return redirect(url_for('admin'))


# ============================================
# МАРШРУТЫ: API ТОВАРОВ
# ============================================
@app.route('/get_product/<int:product_id>')
def get_product(product_id):
    product = next((p for p in load_products() if p['id'] == product_id), None)
    if product:
        return jsonify(product)   # включает image (первое) и images (все)
    return jsonify({'error': 'Товар не найден'}), 404


@app.route('/get_products')
def get_products():
    return jsonify(load_products())


@app.route('/update_product/<int:product_id>', methods=['POST'])
@admin_required
def update_product(product_id):
    data = request.get_json()
    products = load_products()
    for p in products:
        if p['id'] == product_id:
            p['name'] = data.get('name', p['name'])
            p['category'] = data.get('category', p['category'])
            p['description'] = data.get('description', p['description'])
            p['price'] = safe_int(data.get('price', p['price']))
            p['badge'] = data.get('badge', p['badge'])
            break
    save_products(products)
    return jsonify({'success': True})


# ============================================
# МАРШРУТЫ: КОРЗИНА И ИЗБРАННОЕ (СТРАНИЦЫ)
# ============================================
@app.route('/cart')
def cart_page():
    user_id = session.get('user_id')
    is_logged_in = bool(user_id)
    cart = get_user_cart(user_id) if user_id else {}
    return render_template('cart.html',
                           username=session.get('username'),
                           is_logged_in=is_logged_in,
                           initial_cart=cart)


@app.route('/favorites')
def favorites_page():
    user_id = session.get('user_id')
    is_logged_in = bool(user_id)
    favorites = get_user_favorites(user_id) if user_id else []
    return render_template('favorites.html',
                           username=session.get('username'),
                           is_logged_in=is_logged_in,
                           initial_favorites=favorites)


# ============================================
# API: СОХРАНЕНИЕ И ЧТЕНИЕ КОРЗИНЫ / ИЗБРАННОГО
# ============================================
@app.route('/api/cart/save', methods=['POST'])
def api_save_cart():
    user_id = session.get('user_id')
    if not user_id:
        return jsonify({'success': False, 'error': 'Не авторизован'}), 401
    data = request.get_json()
    save_user_cart(user_id, data.get('cart', {}))
    return jsonify({'success': True})


@app.route('/api/favorites/save', methods=['POST'])
def api_save_favorites():
    user_id = session.get('user_id')
    if not user_id:
        return jsonify({'success': False, 'error': 'Не авторизован'}), 401
    data = request.get_json()
    save_user_favorites(user_id, data.get('favorites', []))
    return jsonify({'success': True})


@app.route('/api/cart', methods=['GET'])
def api_get_cart():
    user_id = session.get('user_id')
    if not user_id:
        return jsonify({'cart': {}})
    return jsonify({'cart': get_user_cart(user_id)})


@app.route('/api/favorites', methods=['GET'])
def api_get_favorites():
    user_id = session.get('user_id')
    if not user_id:
        return jsonify({'favorites': []})
    return jsonify({'favorites': get_user_favorites(user_id)})


# ============================================
# ОФОРМЛЕНИЕ ЗАКАЗА
# ============================================
@app.route('/checkout', methods=['POST'])
@login_required
def checkout():
    name = request.form.get('name', '').strip()
    phone = request.form.get('phone', '').strip()
    comment = request.form.get('comment', '').strip()
    cart_data = request.form.get('cart_data', '[]')
    total = safe_int(request.form.get('total', '0'))

    if not name or not phone:
        flash('Заполните имя и телефон', 'error')
        return redirect(url_for('cart_page'))

    if not os.path.exists(ORDERS_FILE):
        init_orders_file()

    wb = load_workbook(ORDERS_FILE)
    ws = wb.active
    max_id = 0
    for row in ws.iter_rows(min_row=2, max_col=1):
        if row[0].value is not None:
            max_id = max(max_id, safe_int(row[0].value))
    order_id = max_id + 1
    order_datetime = datetime.now().strftime('%d.%m.%Y %H:%M')
    user_login = session.get('login', '')

    try:
        items = json.loads(cart_data) if cart_data else []
        items_text = '; '.join([str(item) for item in items])
    except Exception:
        items_text = str(cart_data)

    ws.append([
        order_id, order_datetime, name, phone, user_login,
        items_text, total, comment, 'Новый'
    ])
    last_row = ws.max_row
    for cell in ws[last_row]:
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if cell.column == 7:
            cell.font = Font(bold=True, color='2C1810')
        if cell.column == 9:
            cell.font = Font(bold=True, color='27ae60')
            cell.fill = PatternFill(start_color='E8F8F0', end_color='E8F8F0', fill_type='solid')
    wb.save(ORDERS_FILE)

    # Очищаем корзину после заказа
    user_id = session.get('user_id')
    if user_id:
        save_user_cart(user_id, {})

    return render_template('checkout_success.html',
                           username=session.get('username'),
                           order_id=order_id)


# ============================================
# ПРОСМОТР ЗАКАЗОВ (АДМИНКА)
# ============================================
@app.route('/orders')
@admin_required
def orders_page():
    orders = []
    if os.path.exists(ORDERS_FILE):
        wb = load_workbook(ORDERS_FILE)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0]:
                orders.append({
                    'id': safe_int(row[0]),
                    'datetime': row[1] or '',
                    'client': row[2] or '',
                    'phone': row[3] or '',
                    'login': row[4] if len(row) > 4 else '',
                    'items': row[5] if len(row) > 5 else '',
                    'total': safe_int(row[6]),
                    'comment': row[7] if len(row) > 7 else '',
                    'status': row[8] if len(row) > 8 else 'Новый'
                })
    orders.sort(key=lambda x: x['id'], reverse=True)
    return render_template('orders.html',
                           username=session['username'],
                           orders=orders)


@app.route('/order_details/<int:order_id>')
@admin_required
def order_details(order_id):
    """Детали одного заказа (для модального окна)."""
    if not os.path.exists(ORDERS_FILE):
        return jsonify({'error': 'Файл заказов не найден'}), 404

    wb = load_workbook(ORDERS_FILE)
    ws = wb.active
    order = None
    for row in ws.iter_rows(min_row=2, values_only=True):
        if safe_int(row[0]) == order_id:
            order = {
                'id': safe_int(row[0]),
                'datetime': row[1] or '',
                'client': row[2] or '',
                'phone': row[3] or '',
                'login': row[4] if len(row) > 4 else '',
                'items': row[5] if len(row) > 5 else '',
                'total': safe_int(row[6]),
                'comment': row[7] if len(row) > 7 else '',
                'status': row[8] if len(row) > 8 else 'Новый'
            }
            break
    if not order:
        return jsonify({'error': 'Заказ не найден'}), 404

    # Парсим состав заказа и добавляем информацию о товарах
    products_info = []
    try:
        items_raw = str(order['items'] or '').split('; ')
        for item in items_raw:
            if ':' in item:
                product_id_str, quantity_str = item.split(':')
                product_id = safe_int(product_id_str.strip())
                quantity = safe_int(quantity_str.replace(' шт.', '').strip(), 1)
                product = next((p for p in load_products() if p['id'] == product_id), None)
                if product:
                    products_info.append({
                        'id': product_id,
                        'name': product['name'],
                        'category': product['category'],
                        'price': product['price'],
                        'quantity': quantity,
                        'image': product['image'],   # только первое фото
                        'total': product['price'] * quantity
                    })
    except Exception as e:
        print('Ошибка парсинга состава заказа:', e)

    order['products'] = products_info
    return jsonify(order)


@app.route('/update_order_status', methods=['POST'])
@admin_required
def update_order_status():
    """Изменение статуса заказа."""
    data = request.get_json()
    order_id = data.get('order_id')
    new_status = data.get('status')

    if not order_id or not new_status:
        return jsonify({'success': False, 'error': 'Неверные данные'}), 400
    valid_statuses = ['Новый', 'В обработке', 'Выполнен', 'Отменён']
    if new_status not in valid_statuses:
        return jsonify({'success': False, 'error': 'Недопустимый статус'}), 400
    if not os.path.exists(ORDERS_FILE):
        return jsonify({'success': False, 'error': 'Файл заказов не найден'}), 404

    wb = load_workbook(ORDERS_FILE)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if row[0].value is not None and safe_int(row[0].value) == safe_int(order_id):
            row[8].value = new_status
            text_color, bg_color = STATUS_COLORS.get(new_status, ('000000', 'FFFFFF'))
            row[8].font = Font(bold=True, color=text_color)
            row[8].fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type='solid')
            wb.save(ORDERS_FILE)
            return jsonify({'success': True, 'message': 'Статус обновлён'})

    return jsonify({'success': False, 'error': 'Заказ не найден'}), 404


# ============================================
# ПРОФИЛЬ ПОЛЬЗОВАТЕЛЯ
# ============================================
@app.route('/profile')
@login_required
def profile():
    user = get_user_by_id(session['user_id'])
    if not user:
        return redirect(url_for('logout'))
    return render_template('profile.html', user=user, username=session['username'])


# ============================================
# ОБРАБОТЧИКИ ОШИБОК
# ============================================
@app.errorhandler(404)
def not_found(e):
    return redirect(url_for('index'))


@app.errorhandler(500)
def internal_error(e):
    traceback.print_exc()   # подробный лог в консоль
    return 'Внутренняя ошибка сервера. Подробности в консоли.', 500


# ============================================
# ЗАПУСК
# ============================================
if __name__ == '__main__':
    print("🏛️  Антикварный салон 'Антиквар'")
    print("=" * 50)
    print("🌐 http://localhost:5000")
    print("👤 Админ: admin / admin")
    print("=" * 50)
    app.run(debug=True, host='0.0.0.0', port=5000)